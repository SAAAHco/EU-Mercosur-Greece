# -*- coding: utf-8 -*-
"""
build_workbook_v7_5.py
======================
Regenerates the analytical workbook under the v7.5 specification so that the
DEPOSITED DATA ARTIFACT reproduces the manuscript exactly.

WHAT v7.5 CHANGES
-----------------
1. Tariff-response term (the v7.4 bug): the policy shock is the proportional
   change in the tariff-INCLUSIVE price, (1+tau1)/(1+tau0)-1, the WITS/SMART
   form dM/M = eps*dtau/(1+tau0). v7.4 used -(tau0-tau1)/tau0, reading a 5%
   removal as a 100% price cut (about 21x overstatement per chapter).
2. Incidence-signed wedges baked into Sheet 2 and every projection: exporter-
   borne channels (EUDR, CBAM, EC-2014 compliance) dampen the import response;
   the Greek environmental-tax addon amplifies it (HS 23 net -0.5 pp).

INPUT   : the v7 workbook (counterfactuals, tariff schedules, capacity weights,
          ML results and structural sheets are reused unchanged; only
          projection-derived cells are recomputed).
OUTPUT  : EU_Mercosur_Greece_Analysis_v7_5.xlsx
          plus a printed verification table against the manuscript.

USAGE   : python3 build_workbook_v7_5.py [--src PATH] [--out PATH]
"""
import argparse, random, statistics, copy
import openpyxl
from openpyxl.styles import Font

ap = argparse.ArgumentParser()
ap.add_argument('--src', default='/mnt/project/EU_Mercosur_Greece_Analysis_v7.xlsx')
ap.add_argument('--out', default='/home/claude/work/EU_Mercosur_Greece_Analysis_v7_5.xlsx')
A = ap.parse_args()

EPS_I, EPS_X = -3.5, -2.5
F2F, LAB = 0.10, 0.07
SIGNED = {'23': -0.005, '12': -0.010, '09': -0.010, '47': -0.010,
          '73': -0.030, '76': -0.030, '68': -0.030,
          '38': -0.020, '25': -0.020, '27': -0.005,
          '04': -0.010, '08': -0.010, '15': -0.010, '20': -0.010}
TRQ_CAP_04 = 8_640_000.0

wbv = openpyxl.load_workbook(A.src, data_only=True)   # values for reading
wb  = openpyxl.load_workbook(A.src)                   # target to edit

# ---------------------------------------------------------------- parameters
capw = {}
for r in wbv['13-Per-Chapter Adjustments'].iter_rows(min_row=9, max_row=33, values_only=True):
    capw[str(r[0]).zfill(2)] = float(r[2])

P = {}
for r in wbv['2-Tariff Parameters'].iter_rows(min_row=5, max_row=29, values_only=True):
    hs = str(r[1]).zfill(2)
    P[hs] = dict(t0=r[2] or 0.0, phi=r[4] or 0, m0=r[6] or 0.0, phx=r[7] or 0,
                 w=SIGNED.get(hs, 0.0), capw=capw.get(hs, 0.0))

def eu_tau(hs, y):
    c = P[hs]
    return c['t0'] * max(0.0, 1 - y / c['phi']) if c['t0'] > 0 and c['phi'] else (c['t0'] if c['t0'] else 0.0)

def mer_tau(hs, y):
    c = P[hs]
    return c['m0'] * max(0.0, 1 - y / c['phx']) if c['m0'] > 0 and c['phx'] else (c['m0'] if c['m0'] else 0.0)

def proj_imp(hs, y, cf, wmult=1.0):
    c = P[hs]
    if c['t0'] <= 0 or cf <= 0: return cf
    dprice = (1 + eu_tau(hs, y)) / (1 + c['t0']) - 1          # v7.5 term
    wf = 1 + wmult * c['w'] * min(1.0, y / c['phi']) if c['phi'] else 1.0
    p = cf * (1 + EPS_I * dprice * wf)
    if hs == '04': p = min(p, TRQ_CAP_04)
    return p

def proj_exp(hs, y, cf):
    c = P[hs]
    if c['m0'] <= 0 or cf <= 0: return cf
    dprice = (1 + mer_tau(hs, y)) / (1 + c['m0']) - 1          # v7.5 term
    return cf * (1 + EPS_X * dprice)

def adj_factor(hs, y, f2f=F2F, lab=LAB):
    return 1 - P[hs]['capw'] * (f2f + lab * y / 10)

# ------------------------------------------------- scenario sheets (4 and 5)
def cf_series(sheet):
    """Read CF columns and recover per-chapter growth from the sheet itself."""
    out = {}
    for r in wbv[sheet].iter_rows(min_row=6, max_row=30, values_only=True):
        hs = str(r[1]).zfill(2)
        ci0, ci10 = r[8] or 0.0, r[10] or 0.0
        cx0, cx10 = r[11] or 0.0, r[13] or 0.0
        gi = (ci10 / ci0) ** 0.1 - 1 if ci0 > 0 and ci10 > 0 else 0.0
        gx = (cx10 / cx0) ** 0.1 - 1 if cx0 > 0 and cx10 > 0 else 0.0
        out[hs] = dict(ci0=ci0, gi=gi, cx0=cx0, gx=gx,
                       cfi={0: r[8] or 0.0, 5: r[9] or 0.0, 10: r[10] or 0.0},
                       cfx={0: r[11] or 0.0, 5: r[12] or 0.0, 10: r[13] or 0.0})
    return out

def rewrite_scenario(sheet):
    cf = cf_series(sheet)
    ws = wb[sheet]
    tot = {c: 0.0 for c in range(9, 21)}
    order = []
    for i, r in enumerate(wbv[sheet].iter_rows(min_row=6, max_row=30, values_only=True), start=6):
        hs = str(r[1]).zfill(2); order.append(hs)
        for k, y in ((15, 0), (16, 5), (17, 10)):
            v = proj_imp(hs, y, cf[hs]['cfi'][y]); ws.cell(i, k, v); tot[k] += v
        for k, y in ((18, 0), (19, 5), (20, 10)):
            v = proj_exp(hs, y, cf[hs]['cfx'][y]); ws.cell(i, k, v); tot[k] += v
        for k, y in ((9, 0), (10, 5), (11, 10)):
            tot[k] += cf[hs]['cfi'][y]
        for k, y in ((12, 0), (13, 5), (14, 10)):
            tot[k] += cf[hs]['cfx'][y]
    for c, v in tot.items(): ws.cell(31, c, v)
    # delta block
    hdr_row = next(i for i in range(32, 40)
                   if (wbv[sheet].cell(i, 1).value or '').startswith('Product'))
    ri = hdr_row + 1
    dt = {c: 0.0 for c in (3, 4, 5, 9, 10, 11)}
    for hs in order:
        for (kv, kp, y) in ((3, 6, 0), (4, 7, 5), (5, 8, 10)):
            d = proj_imp(hs, y, cf[hs]['cfi'][y]) - cf[hs]['cfi'][y]
            ws.cell(ri, kv, d); dt[kv] += d
            ws.cell(ri, kp, d / cf[hs]['cfi'][y] if cf[hs]['cfi'][y] else 0.0)
        for (kv, kp, y) in ((9, 12, 0), (10, 13, 5), (11, 14, 10)):
            d = proj_exp(hs, y, cf[hs]['cfx'][y]) - cf[hs]['cfx'][y]
            ws.cell(ri, kv, d); dt[kv] += d
            ws.cell(ri, kp, d / cf[hs]['cfx'][y] if cf[hs]['cfx'][y] else 0.0)
        ri += 1
    if (wbv[sheet].cell(ri, 1).value or '') == 'TOTAL' or ws.cell(ri, 1).value == 'TOTAL':
        for c, v in dt.items(): ws.cell(ri, c, v)
    else:
        ws.cell(ri, 1, 'TOTAL')
        for c, v in dt.items(): ws.cell(ri, c, v)
    return cf

cf1 = rewrite_scenario('4-Scenario S1 (2026)')
cf2 = rewrite_scenario('5-Scenario S2 (2027)')

def totals(cf, y, wmult=1.0, f2f=F2F, lab=LAB):
    TI = TX = TXA = 0.0
    for hs, c in cf.items():
        cfi = c['ci0'] * (1 + c['gi']) ** y
        cfx = c['cx0'] * (1 + c['gx']) ** y
        di = proj_imp(hs, y, cfi, wmult) - cfi
        dx = proj_exp(hs, y, cfx) - cfx
        TI += di; TX += dx; TXA += dx * adj_factor(hs, y, f2f, lab)
    return TI, TX, TXA, TI - TXA

# ------------------------------------------------------------- sheet 6 Summary
ws = wb['6-Summary']
for col, cf, y in ((2, cf1, 0), (3, cf1, 5), (4, cf1, 10), (5, cf2, 5), (6, cf2, 10)):
    CFI = sum(c['ci0'] * (1 + c['gi']) ** y for c in cf.values())
    CFX = sum(c['cx0'] * (1 + c['gx']) ** y for c in cf.values())
    TI, TX, TXA, W = totals(cf, y)
    ws.cell(5, col, CFI); ws.cell(6, col, CFI + TI); ws.cell(7, col, TI)
    ws.cell(8, col, CFX); ws.cell(9, col, CFX + TX); ws.cell(10, col, TX)
    ws.cell(11, col, TI - TX)

# ------------------------------------------------------------- sheet 7 TRQ
ws = wb['7-TRQ Caps Analysis']
p04 = proj_imp('04', 10, cf1['04']['cfi'][10]) / 1e6
for r, slc in ((6, 3.6), (7, 0.72), (8, 4.32)):
    ws.cell(r, 7, p04)
    ws.cell(r, 8, f'Not binding ({slc / p04:,.0f}x headroom)')
p22 = proj_imp('22', 10, cf1['22']['cfi'][10]) / 1e6
ws.cell(12, 7, p22); ws.cell(13, 7, p22)

# ------------------------------------------------------- sheet 8 sensitivity
ws = wb['8-Sensitivity Grid']
MULT = [0.5, 1.0, 1.5, 3.0]; DRAG = [0.07, 0.10, 0.20]
for base, cf in ((9, cf1), (17, cf2)):
    for i, m in enumerate(MULT):
        TI = TX = None
        for j, f in enumerate(DRAG):
            TI, TX, TXA, W = totals(cf, 10, wmult=m, f2f=f)
            ws.cell(base + i, 2 + j, W)
        ws.cell(base + i, 5, TI); ws.cell(base + i, 6, TX)
G = [[totals(cf1, 10, wmult=m, f2f=f)[3] for f in DRAG] for m in MULT]
flat = [v for row in G for v in row]
ws.cell(24, 2, min(flat)); ws.cell(25, 2, G[1][1])
ws.cell(26, 2, max(flat)); ws.cell(27, 2, max(flat) - min(flat))

# ------------------------------------------------------- sheet 9 Monte Carlo
rng = random.Random(42)
acc = {k: {5: [], 10: []} for k in ('imp', 'exp', 'wid')}
for _ in range(1000):
    m = rng.triangular(0.5, 3.0, 1.0)
    f = rng.triangular(0.05, 0.20, 0.10)
    l = rng.triangular(0.03, 0.12, 0.07)
    for y in (5, 10):
        TI, TX, TXA, W = totals(cf1, y, wmult=m, f2f=f, lab=l)
        acc['imp'][y].append(TI); acc['exp'][y].append(TXA); acc['wid'][y].append(W)
def ci(v): s = sorted(v); return statistics.mean(v), s[24], s[974]
ws = wb['9-Monte Carlo CI']
for r, key in ((19, 'imp'), (20, 'exp'), (21, 'wid')):
    m5, l5, h5 = ci(acc[key][5]); m10, l10, h10 = ci(acc[key][10])
    ws.cell(r, 2, m5); ws.cell(r, 3, l5); ws.cell(r, 4, h5)
    ws.cell(r, 5, m10); ws.cell(r, 6, l10); ws.cell(r, 7, h10)
    ws.cell(r, 8, statistics.pstdev(acc[key][10]))

# ------------------------------------------------------ sheet 10 year-by-year
ws = wb['10-Year-by-Year']
order10 = [str(r[1]).zfill(2) for r in wbv['10-Year-by-Year'].iter_rows(min_row=6, max_row=30, values_only=True)]
tots = {sec: [0.0] * 11 for sec in 'ABC'}
for i, hs in enumerate(order10):
    c = cf1[hs]
    for y in range(11):
        cfi = c['ci0'] * (1 + c['gi']) ** y
        cfx = c['cx0'] * (1 + c['gx']) ** y
        di = proj_imp(hs, y, cfi) - cfi
        dxa = (proj_exp(hs, y, cfx) - cfx) * adj_factor(hs, y)
        ws.cell(6 + i, 3 + y, di)
        ws.cell(35 + i, 3 + y, dxa)
        ws.cell(64 + i, 3 + y, di - dxa)
        tots['A'][y] += di; tots['B'][y] += dxa; tots['C'][y] += di - dxa
for base, sec in ((31, 'A'), (60, 'B'), (89, 'C')):
    for y in range(11): ws.cell(base, 3 + y, tots[sec][y])

# --------------------------------------------- sheet 13 capacity adjustments
ws = wb['13-Per-Chapter Adjustments']
tot5 = tot7 = tot8 = 0.0
for i in range(9, 34):
    hs = str(wbv['13-Per-Chapter Adjustments'].cell(i, 1).value).zfill(2)
    c = cf1[hs]
    unadj = proj_exp(hs, 10, c['cfx'][10]) - c['cfx'][10]
    adj = unadj * adj_factor(hs, 10)
    ws.cell(i, 5, unadj); ws.cell(i, 7, adj); ws.cell(i, 8, unadj - adj)
    ws.cell(i, 9, (unadj - adj) / unadj if unadj else 0.0)
    tot5 += unadj; tot7 += adj; tot8 += unadj - adj
ws.cell(34, 5, tot5); ws.cell(34, 7, tot7); ws.cell(34, 8, tot8)
ws.cell(34, 9, tot8 / tot5 if tot5 else 0.0)
ws.cell(39, 9, tot5 * 0.183 * F2F / 1e6)
ws.cell(40, 9, tot8 / 1e6)
ws.cell(41, 9, tot8 / 1e6 - tot5 * 0.183 * F2F / 1e6)

# ---------------------------------------------------- sheet 11 ML vs PE
ws = wb['11-ML vs PE Comparison']
TI5 = totals(cf1, 5)[0]; TI10, TX10, TXA10, W10 = totals(cf1, 10)
rows = [(6, TI5), (7, TI10), (8, TX10), (9, TXA10), (10, W10)]
for r, pe in rows:
    ml = wbv['11-ML vs PE Comparison'].cell(r, 3).value or 0.0
    ws.cell(r, 2, pe); ws.cell(r, 4, pe - ml)
    ws.cell(r, 5, pe / ml if ml else 0.0)
ws.cell(2, 1, 'v7.5: the divergence reflects the ML tariff feature carrying no within-chapter '
              'variation over 2014-2024; the ML output validates the counterfactual baseline '
              'and does not bound the policy effect.')
ws.cell(10, 6, 'PE = $146.0M widening; ML = persistence null (no tariff variation in sample)')

# ---------------------------------------------------- sheet 2 signed wedges
ws = wb['2-Tariff Parameters']
for i in range(5, 30):
    hs = str(wbv['2-Tariff Parameters'].cell(i, 2).value).zfill(2)
    ws.cell(i, 12, SIGNED.get(hs, 0.0))
ws.cell(4, 12, 'Wedge (signed)')

# ------------------------------------------------------------ version sheet
note = wb.create_sheet('0-v7.5 Note', 0)
note.column_dimensions['A'].width = 110
lines = [
 'WORKBOOK v7.5: TARIFF-RESPONSE SPECIFICATION CORRECTION',
 '',
 'The v7.4 pipeline computed the policy shock as the proportional change in the tariff',
 'RATE, -(tau0-tau1)/tau0, reading the removal of a 5% tariff as a 100% buyer-price cut',
 'and overstating the import response by roughly 21x per chapter. v7.5 uses the change',
 'in the tariff-INCLUSIVE price, (1+tau1)/(1+tau0)-1 (WITS/SMART: dM/M = eps*dtau/(1+tau0)),',
 'and bakes the incidence-signed wedges into Sheet 2 and all projections.',
 '',
 'Recomputed sheets: 2 (wedge column), 4, 5, 6, 7, 8, 9 (PE rows 19-21), 10, 11 (PE column), 13.',
 'Unchanged: counterfactual baselines, tariff schedules, capacity weights, ML results',
 '(sheets 1, 3, 9 ML rows, 12, 14) and narrative sheet 15, whose prose may still cite',
 'superseded magnitudes; the quantitative sheets are authoritative.',
 '',
 'Headline mapping (Year 10, $M): widening 3,513.4 -> 146.0; additional imports',
 '4,011.5 -> 197.1; capacity-adjusted exports 498.1 -> 51.1; capacity loss 36.40 -> 3.54.',
 'Code fix: orchestrator_v7_5.py, proj_imports() and proj_exports().',
]
for i, t in enumerate(lines, 1):
    c = note.cell(i, 1, t)
    if i == 1: c.font = Font(bold=True, size=13)

wb.save(A.out)
print('saved', A.out)

# ------------------------------------------------------------ VERIFICATION
print('\n=== VERIFICATION vs MANUSCRIPT ===')
def chk(name, got, want, tol=0.06):
    ok = abs(got - want) <= tol
    print(f'  {"OK " if ok else "FAIL"} {name:34s} workbook {got:>10,.2f}  manuscript {want:>10,.2f}')
    return ok
allok = True
TI10, TX10, TXA10, W10 = totals(cf1, 10)
allok &= chk('additional imports Y10 (M)', TI10/1e6, 197.1)
allok &= chk('adj exports Y10 (M)', TXA10/1e6, 51.1)
allok &= chk('net widening Y10 (M)', W10/1e6, 146.0)
allok &= chk('capacity loss (M)', (TX10-TXA10)/1e6, 3.54, .01)
allok &= chk('grid central (M)', G[1][1]/1e6, 146.0)
allok &= chk('grid span (M)', (max(flat)-min(flat))/1e6, 5.8, .1)
m10, l10, h10 = ci(acc['wid'][10])
allok &= chk('MC mean (M)', m10/1e6, 145.7)
allok &= chk('MC lo (M)', l10/1e6, 143.7)
allok &= chk('MC hi (M)', h10/1e6, 147.8)
W2 = totals(cf2, 10)[3]
allok &= chk('S2 widening (M)', W2/1e6, 161.7)
allok &= chk('delay uplift (%)', 100*(W2/W10-1), 10.8, .1)
allok &= chk('TRQ proj04 (M)', p04, 0.032, .001)
d = {}
for hs, c in cf1.items():
    cfi = c['cfi'][10]; cfx = c['cfx'][10]
    d[hs] = (proj_imp(hs,10,cfi)-cfi) - (proj_exp(hs,10,cfx)-cfx)*adj_factor(hs,10)
allok &= chk('HS23 net (M)', d['23']/1e6, 103.7)
allok &= chk('HS22 net (M)', d['22']/1e6, 36.1)
allok &= chk('HS09 net (M)', d['09']/1e6, 30.7)
allok &= chk('HS73 net (M)', d['73']/1e6, -18.9)
allok &= chk('HS23+09 share (%)', 100*(d['23']+d['09'])/W10, 92.1, .1)
c15 = cf1['15']
n15 = lambda y: (proj_imp('15',y,c15['ci0']*(1+c15['gi'])**y)-c15['ci0']*(1+c15['gi'])**y) - \
                (proj_exp('15',y,c15['cx0']*(1+c15['gx'])**y)-c15['cx0']*(1+c15['gx'])**y)*adj_factor('15',y)
allok &= chk('HS15 net Y8 (M, <0)', n15(8)/1e6, -0.24, .02)
allok &= chk('HS15 net Y9 (M, >0)', n15(9)/1e6, 0.04, .02)
TI3 = totals(cf1,3)[0]; TI5 = totals(cf1,5)[0]; TI7 = totals(cf1,7)[0]
allok &= chk('cadence Y3 (%)', 100*TI3/TI10, 24.4, .6)
allok &= chk('cadence Y5 (%)', 100*TI5/TI10, 45, .6)
allok &= chk('cadence Y7 (%)', 100*TI7/TI10, 70, .6)
print('\nALL CHECKS PASSED' if allok else '\nCHECK FAILURES ABOVE')
