# -*- coding: utf-8 -*-
"""
elasticity_sensitivity.py
=========================
Import-elasticity sensitivity for the Year 10 bilateral widening (Appendix F).

Re-runs the corrected partial-equilibrium projection with the Armington
import elasticity varied across the range supported by the literature
(Yang, Davis & Grant, 2021), holding the export elasticity at -2.5 and all
other parameters at their central values:

    eps_i in {-2.0, -2.5, -3.5 (central), -5.0, -8.0}

The projection uses the corrected tariff shock (proportional change in the
tariff-INCLUSIVE price, the WITS/SMART form)

    pi = (1 + tau_1) / (1 + tau_0) - 1

with incidence-signed non-tariff wedges and the Tier-3 census-anchored
export-capacity adjustment  X * [1 - kappa(c) * (0.10 + 0.07 * y/10)].
The engine below is identical to the one embedded in make_figures_v8.py.

Expected output (manuscript Appendix F):
    Year 10 widening runs from $61.5M at eps = -2.0, through $89.7M at -2.5,
    $146.0M at the central -3.5 and $230.4M at -5.0, to $399.3M at -8.0.

USAGE
-----
    python elasticity_sensitivity.py [--workbook PATH] [--out CSV]

Defaults assume the script is run from the repository root:
    --workbook data/EU_Mercosur_Greece_Analysis_v7_5.xlsx
    --out      elasticity_sensitivity.csv
"""
import argparse
import csv
import openpyxl

# ----------------------------------------------------------------------
# Central calibration (identical to make_figures_v8.py / orchestrator_v7_5.py)
# ----------------------------------------------------------------------
EPS_IMP_CENTRAL, EPS_EXP = -3.5, -2.5
F2F, LAB = 0.10, 0.07
ELASTICITIES = [-2.0, -2.5, -3.5, -5.0, -8.0]

# Incidence-signed wedges, percentage points of per-unit supply cost.
# Negative = borne by the Mercosur exporter, dampens the import response.
# Positive = borne by the Greek producer (env-tax addon inside HS 23 net).
WEDGE = {'23': -0.005,                                   # EUDR -1.0 + env-tax +0.5
         '12': -0.010, '09': -0.010, '47': -0.010,       # EUDR
         '73': -0.030, '76': -0.030, '68': -0.030,       # CBAM tier 1
         '38': -0.020, '25': -0.020,                     # CBAM tier 2
         '27': -0.005,                                   # CBAM tier 3
         '04': -0.010, '08': -0.010, '15': -0.010, '20': -0.010}  # GAP


def load_chapters(path):
    """Read chapter parameters from the analytical workbook (same layout as
    make_figures_v8.py: sheet '13-Per-Chapter Adjustments' for capacity
    weights, sheet '4-Scenario S1 (2026)' rows 6-30 for tariffs and
    counterfactual baselines)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    capw = {}
    for r in wb['13-Per-Chapter Adjustments'].iter_rows(values_only=True):
        if r and r[0] not in (None, 'HS'):
            try:
                capw[str(r[0]).zfill(2)] = float(r[2])
            except (TypeError, ValueError):
                pass
    rows = list(wb['4-Scenario S1 (2026)'].iter_rows(values_only=True))[5:30]
    ch = []
    for r in rows:
        hs = str(r[1]).zfill(2)
        t0, t5 = r[2] or 0, r[3] or 0
        m0, m5 = r[5] or 0, r[6] or 0
        ci0, ci10 = r[8] or 0, r[10] or 0
        cx0, cx10 = r[11] or 0, r[13] or 0
        gi = (ci10 / ci0) ** 0.1 - 1 if ci0 > 0 and ci10 > 0 else 0.0
        gx = (cx10 / cx0) ** 0.1 - 1 if cx0 > 0 and cx10 > 0 else 0.0
        # recover linear phase length from the Year-5 tariff
        phi = 5 / (1 - t5 / t0) if t0 > 0 and t5 / t0 < 1 else (7.0 if t0 > 0 else 0.0)
        phx = 5 / (1 - m5 / m0) if m0 > 0 and m5 / m0 < 1 else (10.0 if m0 > 0 else 0.0)
        ch.append(dict(hs=hs, t0=t0, m0=m0, ci0=ci0, cx0=cx0,
                       gi=gi, gx=gx, phi=phi, phx=phx,
                       w=WEDGE.get(hs, 0.0), capw=capw.get(hs, 0.0)))
    return ch


def project_totals(ch, y, eps_imp):
    """Total additional imports and capacity-adjusted additional exports at
    horizon y, corrected specification. Only the import elasticity varies;
    exports always use EPS_EXP = -2.5."""
    tot_i = tot_x = 0.0
    for c in ch:
        if c['t0'] > 0:
            ty = c['t0'] * max(0.0, 1 - min(y, c['phi']) / c['phi']) if c['phi'] else 0.0
            dprice = (1 + ty) / (1 + c['t0']) - 1            # corrected term
            wf = 1 + c['w'] * min(1.0, y / c['phi']) if c['phi'] else 1.0
            tot_i += c['ci0'] * (1 + c['gi']) ** y * eps_imp * dprice * wf
        if c['m0'] > 0:
            my = c['m0'] * max(0.0, 1 - min(y, c['phx']) / c['phx']) if c['phx'] else 0.0
            dpx = (1 + my) / (1 + c['m0']) - 1               # corrected term
            dx = c['cx0'] * (1 + c['gx']) ** y * EPS_EXP * dpx
            dx *= (1 - c['capw'] * (F2F + LAB * y / 10))
            tot_x += dx
    return tot_i, tot_x


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--workbook', default='data/EU_Mercosur_Greece_Analysis_v7_5.xlsx')
    ap.add_argument('--out', default='elasticity_sensitivity.csv')
    a = ap.parse_args()

    ch = load_chapters(a.workbook)
    print('Import-elasticity sensitivity, Year 10, corrected specification')
    print(f'{"eps_i":>8} {"add. imports $M":>16} {"adj. exports $M":>16} {"widening $M":>13}')
    rows = []
    for eps in ELASTICITIES:
        ti, tx = project_totals(ch, 10, eps)
        tag = '  (central)' if eps == EPS_IMP_CENTRAL else ''
        print(f'{eps:>8.1f} {ti/1e6:>16.1f} {tx/1e6:>16.1f} {ti/1e6 - tx/1e6:>13.1f}{tag}')
        rows.append([eps, round(ti / 1e6, 1), round(tx / 1e6, 1), round((ti - tx) / 1e6, 1)])

    with open(a.out, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['import_elasticity', 'additional_imports_usd_m',
                    'capacity_adjusted_additional_exports_usd_m', 'net_widening_usd_m'])
        w.writerows(rows)
    print(f'\nwritten  {a.out}')


if __name__ == '__main__':
    main()
