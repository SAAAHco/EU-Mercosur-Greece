# -*- coding: utf-8 -*-
"""
make_figures_v8.py
==================
Regenerates the six main-text figures for the JEPP manuscript under the
CORRECTED tariff-response specification.

WHAT CHANGED FROM v7.2
----------------------
orchestrator_v7_4.py line 550 (and 563) computed

    tariff_change = -(initial - cur) / max(initial, 0.001)

which is the proportional change in the TARIFF RATE. Multiplied by the
Armington elasticity it implies that removing a 5% tariff is a 100% price
cut, overstating the import response by roughly 21x.

The correct form is the proportional change in the TARIFF-INCLUSIVE PRICE,
the standard WITS/SMART expression  dM/M = eps * dtau / (1 + tau_0):

    tariff_change = (1 + cur) / (1 + initial) - 1

Wedges are incidence-signed: exporter-borne channels (EUDR, CBAM, GAP)
dampen imports; the Greek environmental-tax addon amplifies them.

OUTPUT
------
PNG (200 dpi, for the Word master) and EPS (vector, for T&F production)
for each of f1_history, f6_pe_ml, f2_chapters, f3_feed, f5_trq,
f4_exposure_capacity.

USAGE
-----
    python3 make_figures_v8.py [--out DIR] [--workbook PATH]
"""
import os, argparse
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
import openpyxl

# ----------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------
ap = argparse.ArgumentParser()
ap.add_argument('--workbook', default='/mnt/project/EU_Mercosur_Greece_Analysis_v7.xlsx')
ap.add_argument('--out', default='/home/claude/work/figs_v8')
A = ap.parse_args()
OUT = A.out
os.makedirs(OUT, exist_ok=True)

# ----------------------------------------------------------------------
# Shared style (matches make_figures_v72.py)
# ----------------------------------------------------------------------
# Okabe-Ito, colour-blind safe. No greys: every element is a saturated hue
# or near-black, so nothing washes out in print or on projection.
C = dict(blue='#0072B2', orange='#E69F00', green='#009E73', verm='#D55E00',
         sky='#56B4E9', purple='#AA4499', navy='#1F3B73', teal='#117777',
         gold='#B8860B', ink='#111111')
plt.rcParams.update({
    'font.family': 'DejaVu Sans', 'font.size': 8.5, 'axes.titlesize': 9,
    'axes.labelsize': 8.5, 'axes.edgecolor': '#111111', 'axes.linewidth': 0.7,
    'axes.spines.top': False, 'axes.spines.right': False,
    'xtick.color': '#111111', 'ytick.color': '#111111',
    'grid.color': '#9FB6CD', 'grid.linewidth': 0.6,
    'figure.dpi': 200, 'savefig.dpi': 200, 'savefig.facecolor': 'white',
    'legend.frameon': False, 'legend.fontsize': 7.6,
    'ps.fonttype': 42, 'pdf.fonttype': 42,
})

def save(fig, name):
    """Write PNG for the Word master and EPS for journal production."""
    fig.savefig(f'{OUT}/{name}.png', bbox_inches='tight')
    fig.savefig(f'{OUT}/{name}.eps', bbox_inches='tight', format='eps')
    plt.close(fig)
    print(f'  written  {name}.png  +  {name}.eps')

def panel(ax, letter):
    ax.set_title(letter, loc='left', fontweight='bold', fontsize=10)

# ----------------------------------------------------------------------
# Engine: corrected partial-equilibrium projection
# ----------------------------------------------------------------------
EPS_IMP, EPS_EXP = -3.5, -2.5
F2F, LAB = 0.10, 0.07

# Incidence-signed wedges, percentage points of per-unit supply cost.
# Negative = borne by the Mercosur exporter, dampens the import response.
# Positive = borne by the Greek producer, amplifies it.
WEDGE = {'23': -0.005,           # EUDR -1.0 + Greek env-tax addon +0.5
         '12': -0.010, '09': -0.010, '47': -0.010,          # EUDR
         '73': -0.030, '76': -0.030, '68': -0.030,          # CBAM tier 1
         '38': -0.020, '25': -0.020,                        # CBAM tier 2
         '27': -0.005,                                      # CBAM tier 3
         '04': -0.010, '08': -0.010, '15': -0.010, '20': -0.010}  # GAP

def load_chapters(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    capw = {}
    for r in wb['13-Per-Chapter Adjustments'].iter_rows(values_only=True):
        if r and r[0] not in (None, 'HS'):
            try: capw[str(r[0]).zfill(2)] = float(r[2])
            except (TypeError, ValueError): pass
    rows = list(wb['4-Scenario S1 (2026)'].iter_rows(values_only=True))[5:30]
    ch = []
    for r in rows:
        hs = str(r[1]).zfill(2)
        t0, t5 = r[2] or 0, r[3] or 0
        m0, m5 = r[5] or 0, r[6] or 0
        ci0, ci10 = r[8] or 0, r[10] or 0
        cx0, cx10 = r[11] or 0, r[13] or 0
        gi = (ci10 / ci0) ** 0.1 - 1 if ci0 > 0 and ci10 > 0 else 0.0
        gx = (cx10 / cx0) ** 0.1 - 1 if cx0 > 0 and cx10 > 0 else 0.0
        # recover linear phase length from the Year-5 tariff
        phi = 5 / (1 - t5 / t0) if t0 > 0 and t5 / t0 < 1 else (7.0 if t0 > 0 else 0.0)
        phx = 5 / (1 - m5 / m0) if m0 > 0 and m5 / m0 < 1 else (10.0 if m0 > 0 else 0.0)
        ch.append(dict(hs=hs, name=str(r[0]), t0=t0, m0=m0, ci0=ci0, cx0=cx0,
                       gi=gi, gx=gx, phi=phi, phx=phx,
                       w=WEDGE.get(hs, 0.0), capw=capw.get(hs, 0.0)))
    return ch, wb

def project(ch, y, wmult=1.0, f2f=F2F, lab=LAB, eps=EPS_IMP, wedge_on=True):
    """Additional imports and capacity-adjusted additional exports at horizon y."""
    det = {}
    for c in ch:
        di = dx = 0.0
        if c['t0'] > 0:
            ty = c['t0'] * max(0.0, 1 - min(y, c['phi']) / c['phi']) if c['phi'] else 0.0
            dprice = (1 + ty) / (1 + c['t0']) - 1          # <-- CORRECTED TERM
            wf = 1 + wmult * c['w'] * min(1.0, y / c['phi']) if (wedge_on and c['phi']) else 1.0
            di = c['ci0'] * (1 + c['gi']) ** y * eps * dprice * wf
        if c['m0'] > 0:
            my = c['m0'] * max(0.0, 1 - min(y, c['phx']) / c['phx']) if c['phx'] else 0.0
            dpx = (1 + my) / (1 + c['m0']) - 1              # <-- CORRECTED TERM
            dx = c['cx0'] * (1 + c['gx']) ** y * EPS_EXP * dpx
            dx *= (1 - c['capw'] * (f2f + lab * y / 10))
        det[c['hs']] = (di, dx)
    return det

CH, WB = load_chapters(A.workbook)
D10 = project(CH, 10)
NET = {h: v[0] - v[1] for h, v in D10.items()}
TOT_I = sum(v[0] for v in D10.values())
TOT_X = sum(v[1] for v in D10.values())
TOT_W = TOT_I - TOT_X
NAME = {c['hs']: c['name'] for c in CH}
print(f'Corrected Year 10: imports {TOT_I/1e6:,.1f}M  adj exports {TOT_X/1e6:,.1f}M  '
      f'widening {TOT_W/1e6:,.1f}M')

# ======================================================================
# Figure 1 - bilateral history, 2014 to 2024 (observed; unaffected by the fix)
# ======================================================================
YRS = np.arange(2014, 2025)
IMP = np.array([384.70, 314.16, 339.56, 323.56, 382.28, 362.96,
                409.99, 521.20, 661.66, 772.31, 697.64])
EXP = np.array([89.95, 74.21, 122.63, 50.95, 179.15, 163.67,
                57.10, 114.89, 88.33, 158.57, 118.33])
BAL = EXP - IMP

fig, ax = plt.subplots(figsize=(7.3, 3.34))
ax.bar(YRS, BAL, color=C['purple'], alpha=.65, width=.68, label='trade balance', zorder=1)
ax.plot(YRS, IMP, '-o', ms=3.6, lw=1.7, color=C['blue'], label='Greek imports from Mercosur', zorder=3)
ax.plot(YRS, EXP, '-o', ms=3.6, lw=1.7, color=C['orange'], label='Greek exports to Mercosur', zorder=3)
ax.axhline(0, color=C['ink'], lw=.7)
ax.set_ylabel('$ million')
ax.set_xticks(YRS)
ax.annotate('2014-2024 CAGR: imports +6.1% per year, exports +2.8%',
            xy=(2014.1, -560), fontsize=7.6, color=C['ink'])
ax.legend(loc='upper left', ncol=1)
ax.grid(axis='y', alpha=.35)
save(fig, 'f1_history')

# ======================================================================
# Figure 2 - the two models at Year 10
# ======================================================================
ML = dict(imp=55.6, exp=39.1, wid=16.6)          # persistence projection
PE = dict(imp=TOT_I / 1e6, exp=TOT_X / 1e6, wid=TOT_W / 1e6)

fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(7.3, 3.15))
labels = ['$\\Delta$ imports', '$\\Delta$ exports\n(adjusted)', 'Widening']
pe = [PE['imp'], PE['exp'], PE['wid']]
ml = [ML['imp'], ML['exp'], ML['wid']]
x = np.arange(3); w = .34
b1 = ax1.bar(x - w/2, pe, w, color=C['blue'], label='Partial equilibrium')
b2 = ax1.bar(x + w/2, ml, w, color=C['orange'], label='Machine learning')
ax1.set_yscale('log'); ax1.set_ylim(8, 600)
ax1.set_xticks(x); ax1.set_xticklabels(labels)
ax1.set_ylabel('Year 10, $ million, log scale')
for b, v in list(zip(b1, pe)) + list(zip(b2, ml)):
    ax1.text(b.get_x() + b.get_width()/2, v * 1.10, f'{v:,.1f}', ha='center', fontsize=7.4)
ax1.grid(axis='y', alpha=.35, which='both')
ax1.legend(loc='upper right')
panel(ax1, 'a')

feats = ['partner code', 'product code', 'ln GDP partner', 'lagged growth',
         'tariff proxy', 'year trend', 'import indicator', 'trade value, lag 2',
         'ln GDP Greece', 'ln distance', 'trade value, lag 1']
vals = [0.7, 1.6, 1.6, 1.6, 2.4, 2.6, 3.3, 4.1, 6.5, 20.0, 55.5]
cols = [C['blue']] * len(vals); cols[feats.index('tariff proxy')] = C['verm']
ax2.barh(feats, vals, color=cols)
for i, v in enumerate(vals):
    ax2.text(v + 1.0, i, f'{v}', va='center', fontsize=7.4)
ax2.set_xlabel('Gain importance, %'); ax2.set_xlim(0, 64)
ax2.text(30, 2.2, 'tariff constant within chapter\nover the estimation window',
         fontsize=7.2, color=C['verm'])
ax2.grid(axis='x', alpha=.35)
panel(ax2, 'b')
fig.tight_layout()
save(fig, 'f6_pe_ml')

# ======================================================================
# Figure 3 - Year 10 net widening by chapter
# ======================================================================
items = sorted(((NET[h] / 1e6, h) for h in NET if abs(NET[h]) > 4e4), reverse=True)
vals = [v for v, _ in items]
CLEAN = {"76": "Aluminum and articles", "23": "Food industries, animal fodder",
         "09": "Coffee, tea, mate, spices", "20": "Prepared vegetables, fruit, olives",
         "15": "Animal and vegetable fats and oils", "73": "Iron and steel articles",
         "08": "Fruit and nuts", "30": "Pharmaceutical products", "22": "Beverages, wine",
         "38": "Chemical products n.e.c.", "84": "Machinery, boilers, reactors",
         "85": "Electrical machinery", "27": "Mineral fuels", "39": "Plastics and articles",
         "82": "Tools and cutlery", "04": "Dairy produce, eggs, honey",
         "68": "Stone and cement articles"}
labs = [f"{h}  {CLEAN.get(h, NAME[h])[:34]}" for _, h in items]
cols = [C['verm'] if v > 0 else C['green'] for v in vals]

fig, ax = plt.subplots(figsize=(7.6, 5.6))
ypos = np.arange(len(vals))[::-1]
ax.barh(ypos, vals, color=cols, height=.68)
ax.set_yticks(ypos); ax.set_yticklabels(labs, fontsize=7.4)
ax.set_xscale('symlog', linthresh=1)
ax.set_xlim(-70, 260)
ax.axvline(0, color=C['ink'], lw=.8)
ax.set_xlabel('Year 10 net widening, $ million (symmetric log scale)')
for yy, v in zip(ypos, vals):
    ax.text(v * 1.30, yy, f'{v:,.1f}',
            va='center', ha='left' if v > 0 else 'right', fontsize=7.2)
ax.text(.98, .06, 'red widens the Greek balance\ngreen narrows it',
        transform=ax.transAxes, ha='right', fontsize=7.4, color=C['ink'])
ax.grid(axis='x', alpha=.35, which='both')
save(fig, 'f2_chapters')

# ======================================================================
# Figure 4 - the composition of exposure in HS 23
# ======================================================================
fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(7.3, 3.25))
crops = ['Wheat', 'Barley', 'Maize', 'Other cereals']
area = [368.7, 93.7, 43.3, 80.3]                      # thousand hectares, ELSTAT 2020
cols = [C['navy'], C['orange'], C['orange'], C['navy']]
ax1.bar(crops, area, color=cols, width=.62)
for i, v in enumerate(area):
    ax1.text(i, v + 9, f'{v:,.0f}', ha='center', fontsize=7.4)
ax1.set_ylabel('thousand hectares')
ax1.set_ylim(0, 430)
ax1.annotate('feed-relevant: 137.0 kha\nagainst 1.87 M livestock units',
             xy=(1.5, 250), fontsize=7.4, color=C['verm'], ha='center')
ax1.grid(axis='y', alpha=.35)
panel(ax1, 'a')

yrs = np.arange(0, 11)
c23 = [c for c in CH if c['hs'] == '23'][0]
series = []
for y in yrs:
    if y == 0:
        series.append(0.0); continue
    ty = c23['t0'] * max(0.0, 1 - min(y, c23['phi']) / c23['phi'])
    dprice = (1 + ty) / (1 + c23['t0']) - 1
    wf = 1 + c23['w'] * min(1.0, y / c23['phi'])
    series.append(c23['ci0'] * (1 + c23['gi']) ** y * EPS_IMP * dprice * wf / 1e6)
ax2.plot(yrs, series, '-o', ms=3.4, lw=1.7, color=C['verm'])
ax2.axvline(c23['phi'], color=C['navy'], ls='--', lw=1.1)
ax2.text(c23['phi'] + .15, max(series) * .35,
         f'phase-out complete\nYear {c23["phi"]:.0f}', fontsize=7.4, color=C['navy'])
ax2.set_xlabel('Year after entry into force')
ax2.set_ylabel('additional HS 23 imports, $ million')
ax2.set_xticks(yrs)
ax2.grid(alpha=.35)
panel(ax2, 'b')
fig.tight_layout()
save(fig, 'f3_feed')

# ======================================================================
# Figure 5 - the eight tariff-rate quotas at Greek scale
# ======================================================================
QUOTA = [('Cheese', 30_000, 5_000), ('Milk powder', 10_000, 3_000),
         ('Honey', 45_000, 4_000), ('Beef', 99_000, 4_500),
         ('Poultry', 180_000, 2_200), ('Sugar', 180_000, 500),
         ('Ethanol (industrial)', 200_000, 700), ('Ethanol (other)', 450_000, 700)]
GR_SHARE = 0.024
d04 = D10['04'][0]
cf04 = [r[10] for r in list(WB['4-Scenario S1 (2026)'].iter_rows(values_only=True))[5:30]
        if str(r[1]).zfill(2) == '04'][0] or 0
proj04 = (cf04 + d04) / 1e6

names = [q[0] for q in QUOTA]
slices = [q[1] * q[2] * GR_SHARE / 1e6 for q in QUOTA]
modelled = [proj04 if n in ('Cheese', 'Milk powder', 'Honey') else np.nan for n in names]

fig, ax = plt.subplots(figsize=(7.3, 3.9))
xp = np.arange(len(names)); w = .36
ax.bar(xp - w/2, slices, w, color=C['blue'], label='notional Greek quota slice (2.4% of EU total)')
ax.bar(xp + w/2, modelled, w, color=C['verm'], label='projected Year 10 Greek imports (HS 04)')
ax.set_yscale('log'); ax.set_ylim(0.008, 90)
ax.set_xticks(xp); ax.set_xticklabels(names, rotation=18, ha='right', fontsize=7.4)
ax.set_ylabel('$ million, log scale')
for i, (s, m) in enumerate(zip(slices, modelled)):
    ax.text(i - w/2, s * 1.14, f'{s:,.2f}', ha='center', fontsize=7.0)
    if not np.isnan(m):
        ax.text(i + w/2, m * 1.14, f'{m:,.3f}', ha='center', fontsize=7.0, color=C['verm'])
        ax.text(i, s * 2.4, f'{s/m:,.0f}x', ha='center', fontsize=8.2,
                fontweight='bold', color=C['ink'])
ax.text(.99, .06, 'headroom shown above each modelled chapter;\nno quota binds at Greek bilateral scale',
        transform=ax.transAxes, ha='right', fontsize=7.4)
ax.legend(loc='upper left', bbox_to_anchor=(0, 1.14), ncol=2)
ax.grid(axis='y', alpha=.35, which='both')
save(fig, 'f5_trq')

# ======================================================================
# Figure 6 - exposure against producer-coalition mobilization capacity
# ======================================================================
CAP = {'23': 'Low', '09': 'Low', '22': 'Medium', '20': 'Medium',
       '15': 'High', '08': 'High', '04': 'High'}
GI = {'Low': 0, 'Medium': 4, 'High': 17}
ORD = {'Low': 0, 'Medium': 1, 'High': 2}

fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(7.3, 2.77))
for h, cap in CAP.items():
    v = NET[h] / 1e6
    ax1.scatter(ORD[cap] + np.random.RandomState(int(h)).uniform(-.12, .12),
                v, s=52, color=C['verm'] if v > 0 else C['green'], zorder=3,
                edgecolor='white', linewidth=.7)
    ax1.annotate(f'{h}', (ORD[cap], v), textcoords='offset points',
                 xytext=(9, 3), fontsize=7.2)
ax1.axhline(0, color=C['ink'], lw=.7)
ax1.set_yscale('symlog', linthresh=1)
ax1.set_xticks([0, 1, 2]); ax1.set_xticklabels(['Low', 'Medium', 'High'])
ax1.set_xlabel('producer-coalition mobilization capacity')
ax1.set_ylabel('Year 10 net widening, $ million')
ax1.set_xlim(-.5, 2.5)
ax1.grid(axis='y', alpha=.35, which='both')
panel(ax1, 'a')

classes = ['Low', 'Medium', 'High']
share = [100 * sum(NET[h] for h in CAP if CAP[h] == k) / TOT_W for k in classes]
xp = np.arange(3); w = .36
ax2b = ax2.twinx()
ax2.bar(xp - w/2, share, w, color=C['verm'], label='share of net exposure')
ax2b.bar(xp + w/2, [GI[k] for k in classes], w, color=C['blue'], label='protected designations')
ax2.axhline(0, color=C['ink'], lw=.7)
ax2.set_xticks(xp); ax2.set_xticklabels(classes)
ax2.set_ylabel('% of net widening', color=C['verm'])
ax2.set_ylim(-22, 108)
ax2b.set_ylabel('protected Greek designations', color=C['blue'])
ax2b.spines['right'].set_visible(True)
for i, v in enumerate(share):
    ax2.text(i - w/2, v + (3 if v > 0 else -9), f'{v:,.1f}%', ha='center', fontsize=7.2)
for i, k in enumerate(classes):
    ax2b.text(i + w/2, GI[k] + .5, f'{GI[k]}', ha='center', fontsize=7.2)
ax2.set_xlabel('mobilization capacity')
panel(ax2, 'b')
fig.tight_layout()
save(fig, 'f4_exposure_capacity')

print(f'\nAll six figures written to {OUT} as PNG and EPS.')
